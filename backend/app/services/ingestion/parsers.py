import csv
import io
import json
from pathlib import Path

import yaml
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


def parse_bytes(file_path: str, data: bytes) -> tuple[str, str]:
    suffix = Path(file_path).suffix.lower()
    if suffix in {".txt", ".md", ".py", ".ts", ".tsx", ".js", ".jsx", ".go", ".rs", ".java", ".tf", ".hcl"}:
        text = data.decode("utf-8", errors="replace")
        return text, "text/plain"
    if suffix == ".json":
        obj = json.loads(data.decode("utf-8"))
        return json.dumps(obj, indent=2), "application/json"
    if suffix in {".yaml", ".yml"}:
        obj = yaml.safe_load(data.decode("utf-8"))
        return yaml.safe_dump(obj), "application/yaml"
    if suffix == ".csv":
        reader = csv.reader(io.StringIO(data.decode("utf-8", errors="replace")))
        rows = [" | ".join(row) for row in reader]
        return "\n".join(rows), "text/csv"
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(pages), "application/pdf"
    if suffix == ".docx":
        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip()), "application/vnd.openxmlformats"
    if suffix == ".xlsx":
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts: list[str] = []
        for sheet in wb.worksheets:
            parts.append(f"# Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                parts.append(" | ".join("" if c is None else str(c) for c in row))
        return "\n".join(parts), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return data.decode("utf-8", errors="replace"), "application/octet-stream"
