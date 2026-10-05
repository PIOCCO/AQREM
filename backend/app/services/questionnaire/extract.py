import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


@dataclass
class ExtractedQuestion:
    external_id: str
    text: str
    section: str | None = None
    expected_answer: str | None = None
    existing_answer: str | None = None


HEADER_ALIASES = {
    "external_id": {"question_id", "id", "ref", "reference", "control_id", "external_id"},
    "text": {"question", "prompt", "text", "requirement", "control"},
    "section": {"section", "category", "domain", "group"},
    "expected_answer": {"expected_answer", "expected", "desired_answer"},
    "existing_answer": {"existing_answer", "current_answer", "response", "answer"},
}


def _match_column(header: str) -> str | None:
    key = header.strip().lower().replace(" ", "_")
    for field, aliases in HEADER_ALIASES.items():
        if key in aliases:
            return field
    return None


def extract_from_csv(data: bytes) -> list[ExtractedQuestion]:
    reader = csv.DictReader(io.StringIO(data.decode("utf-8", errors="replace")))
    if not reader.fieldnames:
        return []
    mapping = {name: _match_column(name) for name in reader.fieldnames if _match_column(name)}
    rows: list[ExtractedQuestion] = []
    for idx, row in enumerate(reader, start=1):
        payload = {mapping[k]: v.strip() for k, v in row.items() if k in mapping and v and v.strip()}
        text = payload.get("text")
        if not text:
            continue
        external_id = payload.get("external_id") or f"Q-{idx:03d}"
        rows.append(
            ExtractedQuestion(
                external_id=external_id,
                text=text,
                section=payload.get("section"),
                expected_answer=payload.get("expected_answer"),
                existing_answer=payload.get("existing_answer"),
            )
        )
    return rows


def extract_from_xlsx(data: bytes) -> list[ExtractedQuestion]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = wb.active
    rows_iter = sheet.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        return []
    headers = [str(c or "").strip() for c in header_row]
    mapping = {i: _match_column(h) for i, h in enumerate(headers) if _match_column(h)}
    extracted: list[ExtractedQuestion] = []
    for idx, row in enumerate(rows_iter, start=1):
        payload: dict[str, str] = {}
        for i, cell in enumerate(row):
            field = mapping.get(i)
            if field and cell is not None and str(cell).strip():
                payload[field] = str(cell).strip()
        text = payload.get("text")
        if not text:
            continue
        external_id = payload.get("external_id") or f"Q-{idx:03d}"
        extracted.append(
            ExtractedQuestion(
                external_id=external_id,
                text=text,
                section=payload.get("section"),
                expected_answer=payload.get("expected_answer"),
                existing_answer=payload.get("existing_answer"),
            )
        )
    return extracted


def extract_from_docx(data: bytes) -> list[ExtractedQuestion]:
    doc = Document(io.BytesIO(data))
    extracted: list[ExtractedQuestion] = []
    section: str | None = None
    counter = 0
    id_pattern = re.compile(r"^([A-Z]{2,5}-\d+)\s*[:\-.]?\s*(.+)$")
    for para in doc.paragraphs:
        line = para.text.strip()
        if not line:
            continue
        if line.isupper() and len(line) < 80:
            section = line.title()
            continue
        match = id_pattern.match(line)
        if match:
            counter += 1
            extracted.append(
                ExtractedQuestion(
                    external_id=match.group(1),
                    text=match.group(2).strip(),
                    section=section,
                )
            )
            continue
        if line.endswith("?"):
            counter += 1
            extracted.append(
                ExtractedQuestion(external_id=f"Q-{counter:03d}", text=line, section=section)
            )
    return extracted


def extract_questionnaire_bytes(filename: str, data: bytes) -> list[ExtractedQuestion]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return extract_from_csv(data)
    if suffix == ".xlsx":
        return extract_from_xlsx(data)
    if suffix == ".docx":
        return extract_from_docx(data)
    if suffix == ".txt":
        lines = [ln.strip() for ln in data.decode("utf-8", errors="replace").splitlines() if ln.strip()]
        out: list[ExtractedQuestion] = []
        for idx, line in enumerate(lines, start=1):
            if line.endswith("?"):
                out.append(ExtractedQuestion(external_id=f"Q-{idx:03d}", text=line))
        return out
    raise ValueError(f"Unsupported questionnaire format: {suffix}")
